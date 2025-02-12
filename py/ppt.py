from pptx import Presentation

# Create presentation object
prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Online Job Finder System"
slide.placeholders[1].text = "PHP MySQL Free Source Code\nYour Name\nDate"

# Slide 2: About Online Job Finder System Project
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "About Online Job Finder System Project"
content = slide.placeholders[1]
content.text = "An innovative platform designed to connect job seekers with employers, providing an efficient and user-friendly job search experience."

# Slide 3: Online Job Finder System Features
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Online Job Finder System Features"
content = slide.placeholders[1]
content.text = ("1. Advanced Search Filters\n"
                "2. Real-Time Job Updates\n"
                "3. Interactive Job Map\n"
                "4. Employer Profiles and Reviews\n"
                "5. Mobile-Friendly Interface")

# Slide 4: Revolutionizing Your Job Search
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Revolutionizing Your Job Search"
content = slide.placeholders[1]
content.text = "Exploring how the Online Job Finder System enhances the job search process through innovative features and tools."

# Slide 5: Understanding the Online Job Finder System
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Understanding the Online Job Finder System"
content = slide.placeholders[1]
content.text = ("1. Advanced Search Filters\n"
                "2. Real-Time Job Updates\n"
                "3. Interactive Job Map\n"
                "4. Employer Profiles and Reviews\n"
                "5. Mobile-Friendly Interface")

# Slide 6: Getting Started with the Online Job Finder System
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Getting Started with the Online Job Finder System"
content = slide.placeholders[1]
content.text = ("Step 1: Create Your Profile\n"
                "Step 2: Explore Job Listings\n"
                "Step 3: Apply to Jobs\n"
                "Step 4: Network and Connect")

# Slide 7: Online Job Finder System Flowchart
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Online Job Finder System Flowchart"
content = slide.placeholders[1]
content.text = "Illustrates the workflow of the system from job posting to application."

# Slide 8: Online Job Finder System ER Diagram
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Online Job Finder System ER Diagram"
content = slide.placeholders[1]
content.text = "Entity-Relationship Diagram showing the database structure and relationships."

# Slide 9: Tools Used
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Tools Used"
content = slide.placeholders[1]
content.text = "PHP, MySQL, HTML, CSS, JavaScript"

# Slide 10: How to Run the Online Job Finder System
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "How to Run the Online Job Finder System"
content = slide.placeholders[1]
content.text = ("1. Download the source code\n"
                "2. Set up the database in MySQL\n"
                "3. Configure PHP settings\n"
                "4. Launch the application in your browser")

# Slide 11: Project Demonstration
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Project Demonstration"
content = slide.placeholders[1]
content.text = "Live demonstration of the Online Job Finder System, showcasing its features and functionalities."

# Slide 12: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conclusion"
content = slide.placeholders[1]
content.text = "The Online Job Finder System streamlines the job search process, making it easier for job seekers to find opportunities and for employers to connect with candidates."

# Save the presentation
ppt_path = './123.pptx'
prs.save(ppt_path)

ppt_path
