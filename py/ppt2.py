from pptx import Presentation
from pptx.util import Inches

# Create a presentation object
prs = Presentation()

# Slide 1: Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Online Job Finder System in PHP MySQL"
subtitle.text = "Your Gateway to Landing the Job of Your Dreams"

# Slide 2: About Online Job Finder System Project
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "About Online Job Finder System Project"
content = slide.placeholders[1]
content.text = (
    "The Online Job Finder System in PHP MySQL Free Source Code is a platform "
    "where users can search for jobs. It provides companies with qualified employees "
    "and allows users to browse jobs based on categories like food, construction, fashion, etc. "
    "Users can view full-time and part-time jobs, sign up, and log in to the system. "
    "The system is developed using PHP and MySQL."
)

# Slide 3: Features
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Online Job Finder System Features"
content = slide.placeholders[1]
content.text = (
    "1. Advanced Search Filters\n"
    "2. Real-Time Job Updates\n"
    "3. Interactive Job Map\n"
    "4. Employer Profiles and Reviews\n"
    "5. Mobile-Friendly Interface"
)

# Slide 4: Advanced Search Filters
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Advanced Search Filters"
content = slide.placeholders[1]
content.text = (
    "Refine your job search based on location, industry, salary range, and job type. "
    "This feature ensures you see opportunities that align with your needs and qualifications."
)

# Slide 5: Real-Time Job Updates
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Real-Time Job Updates"
content = slide.placeholders[1]
content.text = (
    "Stay updated with the latest job listings in real-time. Whether you're looking "
    "for full-time, part-time, or freelance opportunities, our platform keeps you informed."
)

# Slide 6: Interactive Job Map
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Interactive Job Map"
content = slide.placeholders[1]
content.text = (
    "Explore job opportunities geographically using our interactive job map. "
    "Identify potential opportunities in your local area or consider relocation."
)

# Slide 7: Employer Profiles and Reviews
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Employer Profiles and Reviews"
content = slide.placeholders[1]
content.text = (
    "Evaluate potential employers with detailed profiles including company culture, "
    "values, and employee satisfaction ratings. Read and leave reviews to foster trust."
)

# Slide 8: Mobile-Friendly Interface
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Mobile-Friendly Interface"
content = slide.placeholders[1]
content.text = (
    "Search for jobs on the go with our mobile-optimized platform. "
    "Access from your smartphone or tablet with ease, anytime, anywhere."
)

# Slide 9: Getting Started
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Getting Started with the Online Job Finder System"
content = slide.placeholders[1]
content.text = (
    "Step 1: Create Your Profile\n"
    "Step 2: Explore Job Listings\n"
    "Step 3: Apply to Jobs\n"
    "Step 4: Network and Connect"
)

# Slide 10: Tools Used
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Tools Used"
content = slide.placeholders[1]
content.text = (
    "HTML: Page layout and design\n"
    "CSS: Design\n"
    "JavaScript: Frontend\n"
    "PHP: Backend\n"
    "Bootstrap\n"
    "MySQL: Database"
)

# Slide 11: Flowchart Overview
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Online Job Finder System Flowchart"
content = slide.placeholders[1]
content.text = (
    "1. Start\n"
    "2. Sign Up\n"
    "3. Profile Creation\n"
    "4. Explore Job Listings\n"
    "5. Job Matches and Applications\n"
    "6. Application Status and Interview Scheduling\n"
    "7. Interactive Job Map with Location Preferences"
)

# Slide 12: ER Diagram Overview
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Online Job Finder System ER Diagram"
content = slide.placeholders[1]
content.text = (
    "Entities:\n"
    "- USER: UserID, Username, Email, Skills, etc.\n"
    "- JOB: JobID, Title, Description, Location, etc.\n"
    "- APPLICATION: ApplicationID, UserID, JobID, Status, etc.\n"
    "- REVIEW: ReviewID, UserID, Company, Rating, Comment\n\n"
    "Relationships:\n"
    "- USER to APPLICATION: One-to-Many\n"
    "- USER to REVIEW: One-to-Many\n"
    "- JOB to APPLICATION: One-to-Many"
)

# Slide 13: Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conclusion"
content = slide.placeholders[1]
content.text = (
    "The Online Job Finder System simplifies and enhances your job search experience. "
    "With advanced features like real-time updates, interactive maps, and employer reviews, "
    "this system ensures you find the right job quickly and efficiently."
)

# Save the presentation
ppt_path = './ppy.pptx'
prs.save(ppt_path)

ppt_path
